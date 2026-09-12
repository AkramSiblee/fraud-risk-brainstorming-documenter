import argparse
import json
import os

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ARIAL = "Arial"
HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name=ARIAL, size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build(register_rows_path: str, client_name: str, output_path: str):
    with open(register_rows_path) as f:
        rows = json.load(f)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fraud Risk Register"
    ws.sheet_view.showGridLines = False

    ws["B2"] = f"{client_name} — Fraud Risk Register"
    ws["B2"].font = Font(name=ARIAL, size=13, bold=True, color="1F3864")

    headers = ["Risk ID", "Risk Description", "Fraud Triangle Category",
               "Financial Statement Assertion", "Account(s) Affected", "Source", "Planned Response"]
    header_row = 4
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = WRAP

    for r, row in enumerate(rows, start=header_row + 1):
        vals = [row["risk_id"], row["risk_description"], row["fraud_triangle_category"],
                row["assertion"], row["accounts_affected"], row["source"], row.get("planned_response", "")]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BODY_FONT
            cell.alignment = WRAP
            cell.border = BORDER

    widths = [8, 46, 20, 22, 30, 22, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("register_rows_json")
    ap.add_argument("client_name")
    ap.add_argument("-o", "--output", required=True)
    args = ap.parse_args()
    build(args.register_rows_json, args.client_name, args.output)
    print("wrote", args.output)
