import argparse
import os
import sys

from docx import Document
from docx.shared import Pt, RGBColor

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from schema import IngestedPackage
from topics_catalog import topics_in_scope
from framework_terms import terms_for

NAVY = RGBColor(0x1F, 0x38, 0x64)

INDUSTRY_SCHEME_LIBRARY = {
    "construction": ["Percentage-of-completion manipulation", "Change-order timing", "Cost-to-complete estimate bias"],
    "financial institutions": ["Loan-loss reserve manipulation", "Fair value model overrides"],
    "saas": ["Channel stuffing", "Capitalized software cost misclassification"],
    "healthcare": ["Billing/upcoding", "Revenue cycle timing manipulation"],
    "manufacturing": ["Multi-element revenue allocation", "Inventory valuation / shrink concealment", "Bill-and-hold arrangements"],
    "fintech": ["Payment-processing fee timing", "Business email compromise / fraudulent wire transfer"],
}


def _scheme_prompts(industry: str):
    industry_l = (industry or "").lower()
    prompts = []
    for key, schemes in INDUSTRY_SCHEME_LIBRARY.items():
        if key in industry_l:
            prompts.extend(schemes)
    return prompts or ["(No industry-specific scheme library entry matched — use generic scheme prompts.)"]


def build(ingested_json: str, session_type: str, output_path: str):
    ingested = IngestedPackage.from_json(ingested_json)
    md = ingested.metadata
    fw = terms_for(md.primary_framework)

    profile = {
        "complexity_tier": md.complexity_tier,
        "engagement_year": md.engagement_year,
    }
    topics = topics_in_scope(profile)

    doc = Document()
    title = doc.add_paragraph()
    run = title.add_run(f"{md.client_name} — Fraud {fw['discussion_label']} Checklist ({session_type.title()})")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = NAVY

    p = doc.add_paragraph()
    run = p.add_run("Reminder: set aside any prior beliefs about management's honesty and integrity before "
                     f"starting this discussion ({fw['skepticism_citation']}).")
    run.italic = True

    by_category = {}
    for t in topics:
        by_category.setdefault(t.category, []).append(t)

    for category, items in by_category.items():
        h = doc.add_paragraph()
        run = h.add_run(category)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = NAVY
        for t in items:
            doc.add_paragraph(f"\u2610 {t.topic_id} — {t.description}", style=None)

    h = doc.add_paragraph()
    run = h.add_run("Industry-Specific Scheme Prompts")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = NAVY
    for scheme in _scheme_prompts(md.industry):
        doc.add_paragraph(f"\u2610 {scheme}")

    if md.prior_year_findings:
        h = doc.add_paragraph()
        run = h.add_run("Prior-Year Carryforward")
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = NAVY
        doc.add_paragraph(f"\u2610 {md.prior_year_findings} — has this been remediated? Get a specific status, "
                           f"not a repeat of last year's wording.")

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    doc.save(output_path)
    return output_path


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("ingested_json")
    ap.add_argument("session_type", choices=["planning", "interim", "final", "ad-hoc"])
    ap.add_argument("-o", "--output", required=True)
    args = ap.parse_args()
    build(args.ingested_json, args.session_type, args.output)
    print("wrote", args.output)
