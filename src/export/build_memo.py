import argparse
import json
import os
import sys

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from schema import IngestedPackage
from topics_catalog import CATALOG_BY_ID
from framework_terms import terms_for

NAVY = RGBColor(0x1F, 0x38, 0x64)
RED = RGBColor(0xC0, 0x00, 0x00)
GREY = RGBColor(0x59, 0x59, 0x59)

SECTION_TOPICS = [
    ("Framing", ["A1", "A2", "A3"]),
    ("Rebuttable Presumptions", ["B1", "B2"]),
    ("Fraud Triangle Findings", ["C1", "C2", "C3"]),
    ("Prior-Period Continuity", ["D1"]),
    ("Known Fraud & Reporting Channels", ["E1", "E2"]),
    ("IT, Cyber & EUC", ["F1", "F2", "F3"]),
    ("Unpredictability", ["G1"]),
    ("Related Parties & Significant Unusual Transactions", ["H1", "H2"]),
    ("Management Estimates & Accounting Judgments", ["I1"]),
    ("Group / Component Summary", ["J1", "J2", "J3", "J4"]),
    ("First-Year Engagement Considerations", ["K1", "K2"]),
]


def _add_heading(doc, text, size=14, color=NAVY, space_before=12, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p


def _add_body(doc, text, italic=False, bold=False, color=None, size=10.5):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = italic
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return p


def build(ingested_json, topic_coverage_json, flags_json, register_rows_json, output_path):
    ingested = IngestedPackage.from_json(ingested_json)
    with open(topic_coverage_json) as f:
        topics = {t["topic_id"]: t for t in json.load(f)}
    with open(flags_json) as f:
        flags = json.load(f)
    register_rows = []
    if register_rows_json and os.path.exists(register_rows_json):
        with open(register_rows_json) as f:
            register_rows = json.load(f)

    md = ingested.metadata
    fw = terms_for(md.primary_framework)
    doc = Document()

    # ---- Header ----
    if flags["memo_status"] != "Final":
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(flags["memo_status"])
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = RED

    title = doc.add_paragraph()
    run = title.add_run(f"{md.client_name}")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = NAVY

    subtitle = doc.add_paragraph()
    run = subtitle.add_run(f"Fraud Risk Brainstorming Memo — Fiscal Year Ended {md.fiscal_year_end}")
    run.bold = True
    run.font.size = Pt(13)

    _add_body(doc, f"Framework: {md.primary_framework}  |  Complexity tier: {md.complexity_tier}  |  "
                   f"Public-interest entity: {'Yes' if md.public_interest_entity else 'No'}  |  "
                   f"Engagement year: {md.engagement_year}", italic=True, color=GREY, size=9.5)

    _add_heading(doc, f"{fw['discussion_label']}(s) Logged", size=12)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(["Session", "Type", "Date", "Attendees"]):
        hdr[i].text = h
    for s in ingested.sessions:
        row = table.add_row().cells
        row[0].text = s.session_id
        row[1].text = s.session_type + (f" ({s.trigger_source})" if s.trigger_source else "")
        row[2].text = s.discussion_date
        row[3].text = "; ".join(s.attendees)

    _add_body(doc, f"Documentation requirement: {fw['documentation_citation']}.", italic=True, color=GREY, size=9)
    if fw["partner_participation_expected"]:
        _add_body(doc, "This framework expects the engagement partner to participate directly in each "
                       "required session (see attendee lists above).", italic=True, color=GREY, size=9)

    # ---- Topic sections ----
    for section_title, topic_ids in SECTION_TOPICS:
        relevant = [tid for tid in topic_ids if tid in topics]
        if not relevant:
            continue
        _add_heading(doc, section_title)
        for tid in relevant:
            rec = topics[tid]
            catalog_desc = CATALOG_BY_ID[tid].description if tid in CATALOG_BY_ID else ""
            citation = {"B1": fw["b1_citation"], "B2": fw["b2_citation"],
                        "A1": fw["skepticism_citation"]}.get(tid)
            if citation:
                catalog_desc = f"{catalog_desc} ({citation})"
            p = doc.add_paragraph()
            run = p.add_run(f"{tid} — {rec['status']}")
            run.bold = True
            run.font.size = Pt(10.5)
            if rec["status"] == "Not evidenced":
                run.font.color.rgb = RED
            _add_body(doc, catalog_desc, italic=True, color=GREY, size=9)
            if rec.get("supporting_excerpt"):
                src = rec.get("source_session_id") or ""
                _add_body(doc, f"\u201c{rec['supporting_excerpt']}\u201d  — {src}", size=10)
            if rec.get("boilerplate_flag"):
                _add_body(doc, "\u26a0 Flagged: near-identical to the prior-year excerpt for this topic — "
                               "confirm this reflects current-year discussion.", bold=True, color=RED, size=9.5)
            if rec.get("notes_for_reviewer"):
                _add_body(doc, f"\u26a0 {rec['notes_for_reviewer']}", bold=True, color=RED, size=9.5)

    # ---- Fraud Risk Register summary ----
    if register_rows:
        _add_heading(doc, "Fraud Risk Register (summary — see companion workbook for full detail)")
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, h in enumerate(["Risk ID", "Description", "Category", "Accounts Affected"]):
            hdr[i].text = h
        for row in register_rows:
            r = table.add_row().cells
            r[0].text = row["risk_id"]
            r[1].text = row["risk_description"]
            r[2].text = row["fraud_triangle_category"]
            r[3].text = row["accounts_affected"]

    # ---- EQR block ----
    if md.public_interest_entity or any("EQR" in h for h in flags["hard_gates"]):
        _add_heading(doc, "Engagement Quality Review")
        _add_body(doc, "This engagement requires an EQR checkpoint (public-interest entity and/or an open "
                       "Escalation Log entry). EQR completion status: " +
                  ("Logged" if not any("EQR" in h for h in flags["hard_gates"]) else "NOT YET LOGGED"),
                  bold=any("EQR" in h for h in flags["hard_gates"]),
                  color=RED if any("EQR" in h for h in flags["hard_gates"]) else None)

    # ---- Open items ----
    _add_heading(doc, "Open Items (Changelog)")
    if not flags["changelog"] and not flags["hard_gates"]:
        _add_body(doc, "No open items.")
    else:
        for h in flags["hard_gates"]:
            _add_body(doc, f"\u2717 HARD GATE — {h}", bold=True, color=RED, size=10)
        for c in flags["changelog"]:
            _add_body(doc, f"\u2022 {c['item']} — {c['status']}", size=10)
        for g in flags["graduated_flags"]:
            _add_body(doc, f"\u25b8 {g}", size=9.5, color=GREY)

    # ---- Sign-off ----
    _add_heading(doc, "Sign-off")
    table = doc.add_table(rows=3, cols=2)
    table.style = "Light Grid Accent 1"
    table.rows[0].cells[0].text = "Engagement Partner"
    table.rows[0].cells[1].text = "Date"
    table.rows[1].cells[0].text = "Manager"
    table.rows[1].cells[1].text = "Date"
    table.rows[2].cells[0].text = "EQR (if applicable)"
    table.rows[2].cells[1].text = "Date"

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    doc.save(output_path)
    return output_path


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("ingested_json")
    ap.add_argument("topic_coverage_json")
    ap.add_argument("flags_json")
    ap.add_argument("--register", default=None)
    ap.add_argument("-o", "--output", required=True)
    args = ap.parse_args()
    build(args.ingested_json, args.topic_coverage_json, args.flags_json, args.register, args.output)
    print("wrote", args.output)
