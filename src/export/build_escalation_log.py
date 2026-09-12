import argparse
import json
import os
import sys

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from schema import IngestedPackage

ARIAL = "Arial"
HEADER_FILL = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
HEADER_FONT = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name=ARIAL, size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build(ingested_json: str, escalation_fields_json: str, client_name: str, output_path: str):
    ingested = IngestedPackage.from_json(ingested_json)
    with open(escalation_fields_json) as f:
        fields = json.load(f)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Escalation Log"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "C00000"

    ws["B2"] = f"{client_name} — Escalation Log (RESTRICTED ACCESS)"
    ws["B2"].font = Font(name=ARIAL, size=13, bold=True, color="C00000")
    ws["B3"] = "Access limited to the named individuals in each row's Access List. Not part of the general workpaper file until Standard-File Clearance = Yes."
    ws["B3"].font = Font(name=ARIAL, size=9, italic=True, color="595959")

    headers = ["Entry ID", "Description", "Date Identified", "Communication Required",
               "Partner Reviewed", "Access List", "Standard-File Clearance", "Resolution"]
    header_row = 5
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = WRAP

    r = header_row + 1
    for inc in ingested.incidents:
        f_ = fields.get(inc.entry_id, {})
        vals = [
            inc.entry_id, inc.description, inc.date_identified,
            f_.get("communication_required", "Unconfirmed"),
            f_.get("partner_reviewed", "No"),
            f_.get("access_list", ""),
            f_.get("standard_file_clearance", "No"),
            f_.get("resolution", ""),
        ]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BODY_FONT
            cell.alignment = WRAP
            cell.border = BORDER
        r += 1

    widths = [10, 46, 14, 20, 14, 34, 18, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("ingested_json")
    ap.add_argument("escalation_fields_json")
    ap.add_argument("client_name")
    ap.add_argument("-o", "--output", required=True)
    args = ap.parse_args()
    build(args.ingested_json, args.escalation_fields_json, args.client_name, args.output)
    print("wrote", args.output)
